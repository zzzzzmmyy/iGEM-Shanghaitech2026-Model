"""Fluorescence calibration: raw Excel extraction and unchanged M1-M7 evaluation."""
from pathlib import Path
import argparse
import os
import json
import copy
import openpyxl
import pandas as pd
import numpy as np
from sklearn.linear_model import LinearRegression
from sklearn.ensemble import RandomForestRegressor
from sklearn.model_selection import LeaveOneGroupOut
from sklearn.metrics import r2_score, mean_squared_error, mean_absolute_error

def read_grid(path, sheet="Sheet1"):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    grid = {}
    for row in ws.iter_rows():
        for c in row:
            grid[(c.row, c.column)] = c.value
    return grid, ws.max_row

def find_blocks(grid, max_row, row_letters):
    blocks = []
    r = 1
    while r <= max_row:
        v = grid.get((r, 1))
        if isinstance(v, str) and v not in row_letters:
            title = v
            header_row = r + 1
            data_rows = []
            rr = header_row + 1
            while rr <= max_row and grid.get((rr, 1)) in row_letters:
                data_rows.append(rr)
                rr += 1
            blocks.append({"title": title, "data_rows": data_rows})
            r = rr
        else:
            r += 1
    return blocks

def block_to_long(grid, block, value_name):
    recs = []
    for r in block["data_rows"]:
        row_letter = grid[(r, 1)]
        for col in range(2, 14):  
            v = grid.get((r, col))
            if v is None:
                continue
            well_col = col - 1
            recs.append({"row": row_letter, "col": well_col,
                         "well": f"{row_letter}{well_col:02d}", value_name: v})
    return pd.DataFrame(recs)

def parse_file(path, row_letters):
    grid, max_row = read_grid(path)
    blocks = find_blocks(grid, max_row, row_letters)
    assert len(blocks) == 4, f"Expected 4 data blocks, found {len(blocks)}: {[b['title'] for b in blocks]}"
    platemap_b, gfp_b, mch_b, od_b = blocks
    platemap = block_to_long(grid, platemap_b, "label")
    gfp = block_to_long(grid, gfp_b, "GFP_raw")
    mch = block_to_long(grid, mch_b, "mCherry_raw")
    od = block_to_long(grid, od_b, "OD600")

    df = platemap.merge(gfp, on=["row", "col", "well"], how="outer")
    df = df.merge(mch, on=["row", "col", "well"], how="outer")
    df = df.merge(od, on=["row", "col", "well"], how="outer")
    df["label"] = df["label"].astype(str)
    return df, [b["title"] for b in blocks]


def main():
    cli = argparse.ArgumentParser(description=__doc__)
    cli.add_argument("--data-dir", type=Path, default=Path(__file__).resolve().parent,
                     help="Directory containing 20260920.xlsx and 20260921.xlsx")
    cli.add_argument("--out", type=Path, default=Path("output"))
    args = cli.parse_args()
    OUT = str(args.out.resolve())
    os.makedirs(OUT, exist_ok=True)
    for filename, letters, outname in (
        ("20260920.xlsx", "ABDE", "df920_tidy.csv"),
        ("20260921.xlsx", "ABCDEFG", "df921_tidy.csv"),
    ):
        table, titles = parse_file(str(args.data_dir / filename), set(letters))
        table.to_csv(os.path.join(OUT, outname), index=False)
        print(f"{filename}: {len(table)} measured wells")
    df = pd.read_csv(os.path.join(OUT, "df921_tidy.csv"))


    def is_numeric_label(x):
        try:
            float(x); return True
        except (ValueError, TypeError):
            return False


    df["is_frac"] = df["label"].apply(is_numeric_label)
    df.loc[df["is_frac"], "fraction"] = df.loc[df["is_frac"], "label"].astype(float)
    frac_df = df[df["is_frac"]].copy()

    main_grid_vals = {round(0.05 * i, 2) for i in range(21)}
    frac_df["is_extra"] = ~frac_df["fraction"].round(3).isin(main_grid_vals)

    # ---- Background correction using the "blank" (media-only) wells ----
    blank = df[df["label"] == "空白"]
    blank_gfp = blank["GFP_raw"].mean()
    blank_mch = blank["mCherry_raw"].mean()
    print(f"Baseline: GFP channel={blank_gfp:.1f} (SD={blank['GFP_raw'].std(ddof=1):.1f}, n={len(blank)}), "
          f"mCherry channel={blank_mch:.1f} (SD={blank['mCherry_raw'].std(ddof=1):.1f}, n={len(blank)})")

    frac_df["GFPc"] = frac_df["GFP_raw"] - blank_gfp
    frac_df["mCherryc"] = frac_df["mCherry_raw"] - blank_mch

    train = frac_df[~frac_df["is_extra"]].reset_index(drop=True)
    test = frac_df[frac_df["is_extra"]].reset_index(drop=True)
    print(f"Training set (main grid): n={len(train)}, {train['fraction'].nunique()} levels")
    print(f"Test set (extra spike-ins): n={len(test)}, fractions={sorted(test['fraction'].unique())}")

    f0_denom_c = train[train["fraction"] == 0]["GFPc"] + train[train["fraction"] == 0]["mCherryc"]
    print("Post-correction denominator (GFPc+mCherryc) at f=0:", f0_denom_c.values)


    # ============ Candidate model definitions ============
    class RatioModel:
        """Linear fit on the ratio transform: f ~ a*ratio_c + b"""
        def __init__(self):
            self.lr = LinearRegression()

        def fit(self, X, y):
            r = X[:, 0] / (X[:, 0] + X[:, 1])
            self.lr.fit(r.reshape(-1, 1), y)
            return self

        def predict(self, X):
            r = X[:, 0] / (X[:, 0] + X[:, 1])
            return self.lr.predict(r.reshape(-1, 1))


    class RatioPolyModel:
        """Quadratic fit on the ratio transform: f ~ a*ratio_c^2 + b*ratio_c + c"""
        def __init__(self):
            self.lr = LinearRegression()

        def fit(self, X, y):
            r = X[:, 0] / (X[:, 0] + X[:, 1])
            Xp = np.column_stack([r, r**2])
            self.lr.fit(Xp, y)
            return self

        def predict(self, X):
            r = X[:, 0] / (X[:, 0] + X[:, 1])
            Xp = np.column_stack([r, r**2])
            return self.lr.predict(Xp)


    class GFPOnlyLinear:
        def __init__(self):
            self.lr = LinearRegression()

        def fit(self, X, y):
            self.lr.fit(X[:, [0]], y); return self

        def predict(self, X):
            return self.lr.predict(X[:, [0]])


    class MCHOnlyLinear:
        def __init__(self):
            self.lr = LinearRegression()

        def fit(self, X, y):
            self.lr.fit(X[:, [1]], y); return self

        def predict(self, X):
            return self.lr.predict(X[:, [1]])


    class TwoChannelLinear:
        def __init__(self):
            self.lr = LinearRegression()

        def fit(self, X, y):
            self.lr.fit(X, y); return self

        def predict(self, X):
            return self.lr.predict(X)


    class GFPPoly2:
        def __init__(self):
            self.lr = LinearRegression()

        def fit(self, X, y):
            g = X[:, 0]
            Xp = np.column_stack([g, g**2])
            self.lr.fit(Xp, y); return self

        def predict(self, X):
            g = X[:, 0]
            Xp = np.column_stack([g, g**2])
            return self.lr.predict(Xp)


    class RFModel:
        def __init__(self):
            self.rf = RandomForestRegressor(n_estimators=300, max_depth=4, random_state=0, min_samples_leaf=2)

        def fit(self, X, y):
            self.rf.fit(X, y); return self

        def predict(self, X):
            return self.rf.predict(X)


    models = {
        "M1_GFP_only_linear": GFPOnlyLinear(),
        "M2_mCherry_only_linear": MCHOnlyLinear(),
        "M3_two_channel_linear": TwoChannelLinear(),
        "M4_ratio_linear": RatioModel(),
        "M5_ratio_quadratic": RatioPolyModel(),
        "M6_GFP_quadratic": GFPPoly2(),
        "M7_random_forest_two_channel": RFModel(),
    }

    Xtrain = train[["GFPc", "mCherryc"]].values
    ytrain = train["fraction"].values
    groups = train["fraction"].values
    Xtest = test[["GFPc", "mCherryc"]].values
    ytest = test["fraction"].values

    logo = LeaveOneGroupOut()

    cv_results = []
    for name, mdl_proto in models.items():
        oof_pred = np.zeros(len(ytrain))
        for tr_idx, va_idx in logo.split(Xtrain, ytrain, groups):
            m = copy.deepcopy(mdl_proto)
            m.fit(Xtrain[tr_idx], ytrain[tr_idx])
            oof_pred[va_idx] = m.predict(Xtrain[va_idx])
        r2_cv = r2_score(ytrain, oof_pred)
        rmse_cv = np.sqrt(mean_squared_error(ytrain, oof_pred))
        mae_cv = mean_absolute_error(ytrain, oof_pred)

        m_full = copy.deepcopy(mdl_proto)
        m_full.fit(Xtrain, ytrain)
        pred_test = m_full.predict(Xtest)
        r2_test = r2_score(ytest, pred_test)
        rmse_test = np.sqrt(mean_squared_error(ytest, pred_test))
        mae_test = mean_absolute_error(ytest, pred_test)

        cv_results.append(dict(model=name, R2_CV=r2_cv, RMSE_CV=rmse_cv, MAE_CV=mae_cv,
                                R2_test=r2_test, RMSE_test=rmse_test, MAE_test=mae_test))
        print(f"{name}: CV R2={r2_cv:.4f} RMSE={rmse_cv:.4f} MAE={mae_cv:.4f} | "
              f"Test R2={r2_test:.4f} RMSE={rmse_test:.4f} MAE={mae_test:.4f}")

    results_df = pd.DataFrame(cv_results)
    results_df.to_csv(os.path.join(OUT, "model_comparison.csv"), index=False)
    print("\n", results_df.to_string(index=False))

    with open(os.path.join(OUT, "model_blank_baseline.json"), "w") as f:
        json.dump({"blank_gfp": float(blank_gfp), "blank_mch": float(blank_mch)}, f)

    results_df.to_pickle(os.path.join(OUT, "results_df.pkl"))
    train.to_pickle(os.path.join(OUT, "train_921.pkl"))
    test.to_pickle(os.path.join(OUT, "test_921.pkl"))


if __name__ == "__main__":
    main()
